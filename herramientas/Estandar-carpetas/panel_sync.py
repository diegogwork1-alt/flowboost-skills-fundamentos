#!/usr/bin/env python3
"""Hojas de reportes → panel de clientes de Flowboost.

El panel NO lee Meta (no hay META_TOKEN, a propósito): lee lo que ya está en la hoja de
cada cliente. Este script baja cada hoja de `clientes_reportes.json` y envía al panel:

  · `datos`         → Meta por día y campaña
  · `datos-google`  → Google Ads por día y campaña (si la hoja lo tiene)
  · `Reporte Meta`  → las dos casillas azules de cada semana (cierres y facturado) y el fee

Solo LEE las hojas: no escribe nada en Drive. Lo llama `actualizar_todos.py` al terminar;
si falla, el reporte ya está escrito y no pasa nada — se avisa y ya.

Configuración en `panel_sync.json` (junto a este fichero): {"url": "https://flowboost-dashboard.vercel.app"}.
El token NO está ahí: se lee del Llavero de macOS (servicio «panel-flowboost», cuenta INGEST_TOKEN).

Uso:  panel_sync.py [<Cliente> ...] [--dry]
"""
import datetime, json, os, re, sys, tempfile, urllib.request

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


def num(v):
    if v is None or v == "": return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(",", "."))
    except ValueError: return 0.0


def fecha(v):
    if isinstance(v, (datetime.date, datetime.datetime)): return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()[:10]
    return s if re.match(r"\d{4}-\d{2}-\d{2}$", s) else None


def tabla(ws):
    filas = ws.iter_rows(values_only=True)
    cab = [str(c or "").strip() for c in next(filas)]
    for f in filas:
        if any(v not in (None, "") for v in f):
            yield dict(zip(cab, f))


def meta(ws):
    out = []
    for r in tabla(ws):
        f = fecha(r.get("fecha"))
        if not f or not r.get("campana"): continue
        out.append({"fecha": f, "campana": str(r["campana"]), "gasto": round(num(r.get("gasto")), 2),
                    "impresiones": int(num(r.get("impresiones"))), "clics": int(num(r.get("clics"))),
                    "clics_enlace": int(num(r.get("clics_enlace"))), "vistas_landing": int(num(r.get("vistas_landing"))),
                    "alcance": int(num(r.get("alcance"))), "leads": num(r.get("leads"))})
    return out


def google(ws):
    out = []
    for r in tabla(ws):
        f = fecha(r.get("fecha"))
        if not f or not r.get("campana"): continue
        out.append({"fecha": f, "campana": str(r["campana"]), "gasto": round(num(r.get("coste")), 2),
                    "impresiones": int(num(r.get("impresiones"))), "clics": int(num(r.get("clics"))),
                    "clics_enlace": int(num(r.get("clics"))), "leads": num(r.get("conversiones"))})
    return out


def semanas_y_fee(ws):
    """Recorre el reporte: «  SEPTIEMBRE 2026» abre un mes, «Del 7 al 13» es una semana.
    Las columnas se buscan por su TÍTULO, no por su letra: si el diseño cambia, sigue valiendo."""
    sem, anio, mes, col, fee, margen = [], None, None, {}, None, None
    filas = list(ws.iter_rows(values_only=True))
    for i, f in enumerate(filas[:8]):
        for j, v in enumerate(f):
            t = str(v or "")
            abajo = filas[i + 1] if i + 1 < len(filas) else ()
            if t.startswith("Fee de agencia") and j + 1 < len(abajo): fee = num(abajo[j + 1])
            if t.startswith("Margen") and j + 1 < len(abajo): margen = num(abajo[j + 1])
    for f in filas:
        if not f: continue                     # en modo solo lectura una fila vacía llega como ()
        a = str(f[0] or "").strip()
        m = re.match(r"([A-ZÁÉÍÓÚ]+)\s+(\d{4})$", a)
        if m and m.group(1) in MESES:
            mes, anio = MESES.index(m.group(1)) + 1, int(m.group(2)); continue
        if a == "Semana":
            col = {str(v or "").strip(): j for j, v in enumerate(f)}; continue
        m = re.match(r"Del (\d+) al (\d+)$", a)
        if m and mes and col:
            ci = col.get("Cierres (a mano)"); fi = col.get("Facturado (a mano)")
            cierres = f[ci] if ci is not None else None
            fact = f[fi] if fi is not None else None
            sem.append({"desde": f"{anio}-{mes:02d}-{int(m.group(1)):02d}",
                        "hasta": f"{anio}-{mes:02d}-{int(m.group(2)):02d}",
                        "cierres": num(cierres) if cierres not in (None, "") else None,
                        "facturado": num(fact) if fact not in (None, "") else None})
    # Meses con los cierres/facturado apuntados en su fila TOTAL y no en las semanas (Cliente 11):
    # cuenta el total del mes entero, igual que hace el panel al leer la hoja en vivo.
    import calendar
    con_semanas = {s["desde"][:7] for s in sem if s["cierres"] or s["facturado"]}
    mes = anio = None; col = {}
    for f in filas:
        if not f: continue
        a = str(f[0] or "").strip()
        m = re.match(r"([A-ZÁÉÍÓÚ]+)\s+(\d{4})$", a)
        if m and m.group(1) in MESES: mes, anio = MESES.index(m.group(1)) + 1, int(m.group(2)); continue
        if a == "Semana": col = {str(v or "").strip(): j for j, v in enumerate(f)}; continue
        if not (a.upper().startswith("TOTAL") and mes and col): continue
        fu = col.get("Funnel")
        if fu is not None and "TODO" not in str(f[fu] or "").upper(): continue
        clave = f"{anio}-{mes:02d}"
        ci, fi = col.get("Cierres (a mano)"), col.get("Facturado (a mano)")
        c = num(f[ci]) if ci is not None else 0; fa = num(f[fi]) if fi is not None else 0
        if clave in con_semanas or not (c or fa): continue
        sem.append({"desde": f"{clave}-01", "hasta": f"{clave}-{calendar.monthrange(anio, mes)[1]:02d}",
                    "cierres": c or None, "facturado": fa or None})
    return sem, fee, margen


def _limpio(v):
    if isinstance(v, float): return round(v, 6)
    if isinstance(v, (datetime.date, datetime.datetime)): return v.strftime("%Y-%m-%d")
    return v


def tablas_de_reporte(wb):
    """Las pestañas que ve el cliente (`Reporte Meta`, `Reporte Google`, `Histórico hasta …`)
    TAL CUAL las calcula Google Sheets: el panel enseña lo mismo que la hoja, sin recalcular.
    Por pestaña: las casillas de arriba (ticket, margen, cierres, % que cierra, fee) y, por
    mes, sus filas con los valores por NOMBRE de columna (así da igual el orden)."""
    out = []
    for t in wb.sheetnames:
        if not (t.startswith("Reporte") or t.startswith("Histórico")): continue
        filas = [f for f in wb[t].iter_rows(values_only=True)]
        arriba = {}
        for i, f in enumerate(filas[:8]):
            abajo = filas[i + 1] if i + 1 < len(filas) else ()
            for j, v in enumerate(f or ()):
                txt = str(v or "").strip()
                if txt and len(txt) < 60 and not txt.startswith("PARA ") and j + 1 < len(abajo) \
                        and isinstance(abajo[j + 1], (int, float)):
                    arriba[re.sub(r"\s*\(.*?\)\s*$", "", txt)] = _limpio(abajo[j + 1])
        meses, mes, cols = [], None, None
        for f in filas:
            if not f or f[0] in (None, ""): continue
            a = str(f[0]).strip()
            m = re.match(r"([A-Za-zÁÉÍÓÚáéíóú]+)\s+(\d{4})$", a)
            if m and m.group(1).upper() in MESES:
                mes = {"mes": MESES.index(m.group(1).upper()) + 1, "anio": int(m.group(2)), "filas": []}
                meses.append(mes); cols = None; continue
            if a == "Semana":
                cols = [str(v or "").strip() for v in f]; continue
            if mes is None or not cols: continue
            tipo = "total" if a.upper().startswith("TOTAL") else "semana" if a.startswith("Del ") else "otra"
            if tipo == "otra" and len(a) > 40: continue      # notas de la hoja, no filas de la tabla
            mes["filas"].append({"tipo": tipo, "valores": {c: _limpio(v) for c, v in zip(cols, f) if c}})
        out.append({"pestana": t, "arriba": arriba,
                    "columnas": [c for c in (cols or []) if c],
                    "meses": [m for m in meses if m["filas"]]})
    return out


def enviar(cfg, cuerpo):
    req = urllib.request.Request(cfg["url"].rstrip("/") + "/api/ingest", method="POST",
        data=json.dumps(cuerpo).encode(), headers={"content-type": "application/json",
                                                   "x-token": cfg["token"], "user-agent": "panel_sync"})
    return json.load(urllib.request.urlopen(req, timeout=60))


def main():
    import actualizar_datos as ad
    from actualizar_todos import leer_clientes
    from openpyxl import load_workbook
    dry = "--dry" in sys.argv
    solo = {a.lower() for a in sys.argv[1:] if not a.startswith("--")}
    p = os.path.join(AQUI, "panel_sync.json")
    if not os.path.exists(p) and not dry:
        print("· panel: sin panel_sync.json, no se envía nada"); return 0
    cfg = json.load(open(p)) if os.path.exists(p) else {}
    if not cfg.get("token"):                      # el token vive en el Llavero de macOS
        import subprocess
        r = subprocess.run(["security", "find-generic-password", "-s", "panel-flowboost", "-a", "INGEST_TOKEN", "-w"],
                           capture_output=True, text=True)
        cfg["token"] = r.stdout.strip()
        if not cfg["token"]: print("✗ panel: no encuentro INGEST_TOKEN en el Llavero"); return 1
    fallos = 0
    for c in leer_clientes(solo_activos=False):
        if not c.get("sheet_id") or (solo and c["cliente"].lower() not in solo): continue
        try:
            tmp = os.path.join(tempfile.mkdtemp(), "h.xlsx")
            ad.bajar(c["sheet_id"], tmp)
            wb = load_workbook(tmp, data_only=True, read_only=True)
            cuerpo = {"cliente": c["cliente"], "sheet_id": c["sheet_id"]}
            if "datos" in wb.sheetnames: cuerpo["meta"] = meta(wb["datos"])
            if "datos-google" in wb.sheetnames: cuerpo["google"] = google(wb["datos-google"])
            if "Reporte Meta" in wb.sheetnames:
                cuerpo["semanas"], cuerpo["fee"], cuerpo["margen"] = semanas_y_fee(wb["Reporte Meta"])
            cuerpo["reportes"] = tablas_de_reporte(wb)
            resumen = (f"{len(cuerpo.get('meta', []))} Meta · {len(cuerpo.get('google', []))} Google · "
                       f"{sum(1 for s in cuerpo.get('semanas', []) if s['cierres'] or s['facturado'])} semanas con cierres")
            if dry: print(f"  {c['cliente']}: {resumen} (--dry, no se envía)"); continue
            enviar(cfg, cuerpo)
            print(f"✓ panel · {c['cliente']}: {resumen}")
        except Exception as e:
            fallos += 1
            print(f"✗ panel · {c['cliente']}: {e}")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
